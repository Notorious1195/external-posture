"""HTTPS/TLS sites and applications blocked by a WAF: paths the testaspnet (HTTP-only) run never exercised."""
import json
from io import BytesIO
from types import SimpleNamespace
from docx import Document
from access_policy import normalize
from intelligence import build_observations
from posture import calculate,check,losses,explanation
from finding_model import verification
from report_ai import compose,validate_prose,incomplete_tools
from docx_report import report_docx

URL='https://example.com'
TLS_XML=('<nmaprun><script id="ssl-enum-ciphers" output=""><table key="TLSv1.0"><table key="ciphers"><table><elem key="name">TLS_RSA_WITH_AES_128_CBC_SHA</elem>'
         '<elem key="strength">A</elem></table></table></table><table key="TLSv1.2"><table key="ciphers"><table><elem key="name">TLS_ECDHE_RSA_WITH_AES_128_GCM_SHA256</elem>'
         '<elem key="strength">A</elem></table></table></table></script></nmaprun>')

def base(**extra):
    d={'domain':'example.com','organization':'Empresa de prueba','project_name':'Evaluación externa','date':'2026-10-01','hosts':['example.com'],'scanned':['example.com'],
       'ports':[{'host':'example.com','ip':'93.184.216.34','port':'443','protocol':'tcp','service':'https'}],'subdomains':['example.com'],'resolved':['example.com'],
       'explicit_ips':[],'dns_ips':{},'visual_evidence':[],'historical_count':0,
       'results':[{'tool':'HTTPx Toolkit','status':'OK','output':URL+' [200]'},{'tool':'Nmap web/TLS · '+URL+' · 93.184.216.34','status':'OK','output':TLS_XML}]}
    d.update(extra);normalize(d)
    d['observations'],d['coverage']=build_observations(d)
    return d

def https_site():
    return base(posture_probe=[check('W-01',URL,'Evaluado',0,'x-frame-options: DENY'),check('W-02',URL,'Evaluado',0,'x-content-type-options: nosniff'),
                               check('T-03',URL,'Evaluado',0,'Redirección a HTTPS'),check('T-01',URL,'Evaluado',.6,'certificate has expired')])

def by_control(d,control):return next(o for o in d['observations'] if o.get('control')==control)

def test_tls_findings_risk_and_verification():
    d=https_site()
    t01,t02=by_control(d,'T-01'),by_control(d,'T-02')
    assert (t01['risk'],t02['risk'])==('Medio','Medio')
    assert '"TLSv1.0"' in t02['fact']
    assert not any(o.get('control')=='T-03' for o in d['observations']),'HTTPS with redirect must not produce a plaintext finding'
    assert 'ssl-enum-ciphers' in verification(t02) and 'vigencia' in verification(t01)

def test_tls_losses_and_summary():
    d=https_site();p=calculate(d);l=losses(p)
    assert round(l['Certificado HTTPS verificable'],2)==round(l['Protocolos y cifrados TLS'],2)==3.0
    assert 'Transporte web cifrado' not in l
    assert 'certificado HTTPS verificable' in explanation(d,p)

def test_tls_prose_checks():
    t02=by_control(https_site(),'T-02')
    assert validate_prose('El servidor negocia TLSv1.0 además de TLSv1.2, un protocolo antiguo que conviene retirar del servicio publicado.',t02)
    try:validate_prose('El servidor 10.0.0.5 negocia TLSv1.0 además de TLSv1.2 en el servicio publicado.',t02);raise AssertionError('IP no observada aceptada')
    except ValueError as exc:assert '10.0.0.5' in str(exc)

def test_https_report_renders_specific_checks():
    d=https_site();compose(d,None,SimpleNamespace(write=lambda *a:None))
    text='\n'.join(p.text for p in Document(BytesIO(report_docx(d))).paragraphs)
    assert 'Repetir la enumeración de protocolos' in text and 'Contenido servido por HTTP sin cifrado' not in text

def blocked_site():
    alert={'pluginId':'10202','alertRef':'10202','url':URL+'/','risk':'Medium','confidence':'Medium','name':'Absence of Anti-CSRF Tokens','description':'d','solution':'s','evidence':'<form>'}
    return base(web_preflight=[{'url':URL,'status':'OK','http_status':200,'blocked':True,'reason':'Acceso restringido o página de protección'}],
                posture_probe=[check('T-01',URL,'Evaluado',0,'Certificado válido')],
                zap_reports=[{'url':URL,'status':'OK','crawler_complete':True,'passive_pending':0,'alerts':[alert]}],
                results=[{'tool':'HTTPx Toolkit','status':'OK','output':URL+' [200]'},{'tool':'Nmap web/TLS · '+URL+' · 93.184.216.34','status':'OK','output':TLS_XML},
                         {'tool':'Acceso antes de pruebas · '+URL,'status':'OK','output':'{}'},{'tool':'Nikto · '+URL,'status':'PARCIAL','output':''}])

def test_waf_block_keeps_transport_and_drops_web_attribution():
    d=blocked_site();p=calculate(d)
    assert not any(o.get('zap_rule') for o in d['observations']),'ZAP alerts behind a block page are not the application'
    web=[c for c in p['checks'] if c['dimension'] in ('V','W')]
    assert web and all(c['status']=='No evaluado' for c in web)
    assert {c['control']:c['status'] for c in p['checks'] if c['dimension']=='T'}['T-02']=='Evaluado','transport stays assessable'
    assert p['category']=='No concluyente'

def test_waf_block_report_states_limits():
    d=blocked_site();compose(d,None,SimpleNamespace(write=lambda *a:None))
    assert 'Nikto · https://example.com (PARCIAL)' in incomplete_tools(d)
    assert 'Nikto' in d['narrative']['limitations'] and d['report_ai']['status']=='REGLAS'
    text='\n'.join(p.text for p in Document(BytesIO(report_docx(d))).paragraphs)
    assert 'No concluyente' in '\n'.join(c.text for t in Document(BytesIO(report_docx(d))).tables for r in t.rows for c in r.cells)
    assert 'No se utilizó texto de IA' in text
