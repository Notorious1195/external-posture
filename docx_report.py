"""Readable evidence-led Word report for external exposure assessments."""
import base64
import hashlib
import re
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from consulting_report import visual_for
from intelligence import narrative
from assets import dashboard_data


def missing_visuals(data):
    return [item for item in data.get("observations", []) if not visual_for(data, item)]


def _verified_image(visual):
    raw = base64.b64decode(visual["image_b64"], validate=True)
    if len(raw) > 3_000_000 or hashlib.sha256(raw).hexdigest() != visual["sha256"]:
        raise ValueError("La imagen no coincide con su registro de integridad: " + visual.get("id", ""))
    image = Image.open(BytesIO(raw))
    if image.format not in ("PNG", "JPEG"):
        raise ValueError("Formato de evidencia no admitido")
    image.verify()
    return raw


def _p(doc, label, text):
    p = doc.add_paragraph()
    p.add_run(label + " ").bold = True
    p.add_run(str(text))
    return p



SHOWN_URLS=4

def report_docx(data):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from posture import calculate,es_num
    from finding_model import verification
    from finding_model import human_description
    from report_ai import approved_note
    from redaction import text as redact
    from intel_context import context as intelligence_context
    from report_context import executive,priorities
    from nmap_checks import tls_summary
    from collections import Counter
    import json
    doc=Document();sec=doc.sections[0]
    sec.page_width,sec.page_height=Inches(8.5),Inches(11)
    sec.left_margin=sec.right_margin=Inches(.8);sec.top_margin=sec.bottom_margin=Inches(.7)
    styles=doc.styles
    styles['Normal'].font.name='Calibri';styles['Normal'].font.size=Pt(11)
    styles['Normal'].paragraph_format.space_after=Pt(7)
    styles['Normal'].paragraph_format.line_spacing=1.08
    for name,size,color in [('Title',28,'122F4B'),('Heading 1',18,'122F4B'),('Heading 2',14,'087E8B'),('Heading 3',12,'122F4B')]:
        st=styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color)
    for style in styles:
        for borders in list(style.element.iter(qn('w:pBdr'))):borders.getparent().remove(borders)
    BLUE='122F4B';TEAL='087E8B';LIGHT='EEF5F8';WHITE='FFFFFF'
    palette={'Crítico':'B42318','Alto':'C75A12','Medio':'9A6700','Bajo':'167348','Informativo':'32658A'}
    p=calculate(data);cov=data.get('coverage',{});n=narrative(data)
    org=data.get('organization') or data['domain'];project=data.get('project_name') or 'Evaluación de postura externa'
    def shade(cell,color):
        el=OxmlElement('w:shd');el.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(el)
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for cell,label in zip(t.rows[0].cells,headers):
            cell.text=str(label);shade(cell,BLUE)
            for run in cell.paragraphs[0].runs:run.font.color.rgb=RGBColor.from_string(WHITE);run.bold=True
        for index,row in enumerate(rows):
            cells=t.add_row().cells
            for c,value in zip(cells,row):c.text=redact(str(value));shade(c,LIGHT if index%2==0 else WHITE)
        return t
    def paragraph(value):
        for line in str(value).split('\n'):
            if line.strip():doc.add_paragraph(redact(line.strip()))
    def label(name,value):_p(doc,name,redact(str(value)))
    def risk_label(item):
        line=doc.add_paragraph();run=line.add_run(item.get('risk','Informativo').upper());run.bold=True;run.font.color.rgb=RGBColor.from_string(palette.get(item.get('risk'),BLUE))
        line.add_run('  '+item.get('risk_reason',''))
    header=sec.header.paragraphs[0]
    header.add_run(org+'  |  POSTURA EXTERNA').font.color.rgb=RGBColor.from_string(TEAL)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Postura externa  ·  '+str(data.get('date',''))+' UTC  ·  Página ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    # Cover and document controls.
    doc.add_paragraph('Informe de postura de seguridad externa',style='Title')
    subtitle=doc.add_paragraph(org);subtitle.runs[0].font.size=Pt(22);subtitle.runs[0].font.color.rgb=RGBColor.from_string(TEAL)
    paragraph(project)
    label('Dominio y alcance principal',data['domain']);label('Fecha de evaluación UTC',data.get('date',''))
    score=table(['Cumplimiento acreditado','Clasificación','Cobertura de controles'],[[es_num(p['score'])+'/100' if p['score'] is not None else 'Sin datos',p['category'],es_num(p['coverage'])+'%']])
    for c in score.rows[1].cells:
        for run in c.paragraphs[0].runs:run.bold=True;run.font.size=Pt(18);run.font.color.rgb=RGBColor.from_string(TEAL)
    paragraph('Evaluación de exposición y controles visibles desde Internet mediante reconocimiento y ciberinteligencia de fuentes públicas. El informe reúne conclusiones, prioridades, evidencias y acciones de mejora sobre los activos evaluados.')
    table(['Control del documento','Valor'],[
      ['Versión de aplicación y metodología',p['version']],['Responsable del informe',data.get('report_author') or 'Por asignar por el propietario'],
      ['Distribución',data.get('distribution') or 'Por definir por el propietario'],['Estado documental','Preparado automáticamente; revisión de distribución a cargo del propietario'],
      ['Autorización y alcance',data.get('authorization_note') or 'Alcance declarado y confirmado por el operador en la aplicación'],
      ['Punto de observación',data.get('scanner_public_ip') or 'Kali configurado por el operador; IP pública de salida no registrada']])
    paragraph('La puntuación procede de reglas reproducibles. '+('Se incorporó redacción asistida con controles automáticos acotados; requiere revisión profesional antes de su distribución.' if data.get('report_ai',{}).get('reviewed',0) else 'No se utilizó texto de IA en esta versión del documento; la redacción procede del catálogo y de los datos registrados.'))
    doc.add_page_break()
    doc.add_heading('Contenido del informe',1)
    for title in ['Resumen ejecutivo','Alcance y metodología','Lectura de los resultados','Postura de seguridad externa','Hallazgos y observaciones con evidencia','Plan de mejora prioritario','Ciberinteligencia y contexto público','Anexos de cobertura y trazabilidad']:
        doc.add_paragraph(title,style='List Bullet')
    doc.add_heading('1 Resumen ejecutivo',1)
    paragraph(data.get('narrative',{}).get('executive') or executive(data))
    top=priorities(data)
    if top:
        doc.add_heading('Prioridades para la organización',2)
        table(['Prioridad','Condición','Riesgo'],[[str(i),o['title'],o['risk']] for i,o in enumerate(top[:5],1)])
    doc.add_heading('2 Alcance y metodología',1)
    paragraph('La evaluación construye un inventario de nombres e IPs públicas mediante fuentes de descubrimiento, comprueba su resolución y evalúa todos los objetivos elegibles conocidos. Se revisan servicios, aplicaciones, respuestas HTTP, transporte cifrado, correo/DNS y señales de fuentes públicas. ZAP recorre aplicaciones y analiza respuestas de forma pasiva; no se presenta como una prueba exhaustiva de explotación.')
    inv=dashboard_data(data)
    table(['Indicador','Resultado'],[[k,v] for k,v in [('Nombres candidatos',cov.get('candidates',0)),('Nombres con DNS A',cov.get('dns_a',0)),('IPs observadas en DNS',inv['dns_ips']),('IPs declaradas sin FQDN observado',inv['ip_without_fqdn']),('Objetivos elegibles conocidos',cov.get('selected',0)),('Objetivos con Nmap completado',cov.get('scanned',0)),('Puertos abiertos observados',cov.get('open_ports',0)),('URLs históricas devueltas',cov.get('historical_urls',0))]])
    paragraph('Comprobación de mejoras: cada ficha indica la comprobación concreta que acreditaría la corrección. Debe repetirse desde un punto externo autorizado, conservando la respuesta fechada y una captura, y verificando el funcionamiento legítimo después del cambio. Un bloqueo de la herramienta o una respuesta incompleta no acredita corrección.')
    paragraph('Una IP relacionada por DNS puede pertenecer a una CDN o a un proveedor compartido; esa relación no demuestra titularidad de la organización ni autoriza ampliar el escaneo a su red. Las respuestas dependen del momento y del punto de observación. Cero resultados en una fuente histórica no demuestra que nunca existieran páginas.')
    doc.add_heading('3 Lectura de los resultados',1)
    paragraph(n['assessment'])
    counts=Counter(o.get('risk','Informativo') for o in data.get('observations',[]))
    table(['Tipo de prioridad','Condiciones'],[[risk,counts[risk]] for risk in ('Crítico','Alto','Medio','Bajo','Informativo')])
    paragraph('Las observaciones informativas describen contexto y puntos de revisión sin atribuir una debilidad. Las descripciones distinguen condiciones comprobadas de señales pendientes de corroboración. No se duplica una condición porque varias herramientas la detecten, ni se presume explotación a partir de una alerta pasiva.')
    doc.add_heading('4 Postura de seguridad externa',1)
    table(['Dimensión','Peso','Nota','Cobertura'],[[d['name'],str(round(d['weight']*100))+'%',es_num(d['score'])+'/100' if d['score'] is not None else 'No aplica',str(d['done'])+'/'+str(d['planned'])] for d in p['dimensions']])
    paragraph('Cada nota refleja cumplimiento acreditado dentro de los controles previstos. Un control pendiente aporta cero y conserva su condición de no evaluado; un control excluido necesita justificación. Un valor de 100 en servicios significa que no se observaron los servicios sensibles definidos entre los puertos examinados, no que toda la infraestructura sea segura.')
    label('Rango por controles pendientes',es_num(p['interval'][0])+' – '+es_num(p['interval'][1]));label('Metodología',p['version']+'; experimental, pendiente de calibración empírica')
    for d in p['dimensions']:
        doc.add_heading(d['name'],2)
        checks=[c for c in p['checks'] if c['dimension']==d['dimension'] and c['status']!='No aplica']
        bad=[c for c in checks if c['status']=='Evaluado' and c.get('penalty',0)>0];unknown=[c for c in checks if c['status']=='No evaluado']
        paragraph(f"Se completaron {d['done']} de {d['planned']} comprobaciones aplicables, con {d['passed']} plenamente cumplidas. "+('Reducen la nota: '+', '.join(dict.fromkeys(c['title'] for c in bad))+'. ' if bad else 'Las comprobaciones completadas no registran penalizaciones en esta dimensión. ')+('Quedan pendientes: '+', '.join(dict.fromkeys(c['title'] for c in unknown))+'.' if unknown else 'No hay controles pendientes en esta dimensión.'))
    doc.add_heading('5 Hallazgos y observaciones con evidencia',1)
    paragraph('Cada ficha reúne una condición por aplicación o activo, sus ubicaciones y las fuentes que la sustentan. La prioridad cualitativa es independiente de la ecuación global de postura. Los identificadores HAL permanecen estables entre ejecuciones mientras se mantenga la misma condición y origen.')
    for item in data.get('observations',[]):
        doc.add_heading(item['id']+'  '+item['title'],2)
        risk_label(item);label('Activo',item['asset'])
        description,meaning,remediation=human_description(item)
        note=approved_note(data,item)
        doc.add_heading('Descripción y relevancia',3)
        paragraph(description)
        if item.get('status')!='Confirmado':paragraph('La condición procede de una señal automática. Su descripción documenta lo observado; aún debe verificarse su aplicabilidad o impacto antes de afirmar una vulnerabilidad explotable.')
        paragraph(note['explanation'] if note else meaning)
        if item.get('triage'):
            label('Parámetros prioritarios para revisión',', '.join(item['triage']['prioritize']) or 'Sin parámetro funcional identificado')
            label('Campos ocultos que conservan revisión contextual',', '.join(item['triage']['hidden_fields']))
            paragraph(item['triage']['note'])
        urls=item.get('affected_urls',[item['asset']])
        # Long route lists move to the annex; the finding keeps a few and a cross-reference.
        label('Ubicaciones afectadas','\n'.join(urls[:SHOWN_URLS])+(f'\n… y {len(urls)-SHOWN_URLS} más (anexo «Ubicaciones afectadas por hallazgo», {item["id"]})' if len(urls)>SHOWN_URLS else ''))
        doc.add_heading('Evidencia y trazabilidad',3)
        label('Referencias técnicas',', '.join(item.get('evidence',[])))
        visuals=visual_for(data,item)
        for visual in visuals:
            raw=_verified_image(visual);im=Image.open(BytesIO(raw));w,h=im.size
            width=min(6.85,6.5*w/max(h,1))
            paragraph('Evidencia '+visual['id']+' · '+visual.get('source',''))
            doc.add_picture(BytesIO(raw),width=Inches(width))
            label('Captura UTC',visual.get('captured_at_utc',''));label('SHA 256',visual['sha256'])
        if not visuals:paragraph('La evidencia técnica está vinculada en el registro de esta ejecución. El documento conserva condición de borrador mientras falte su soporte visual.')
        doc.add_heading('Remediación recomendada',3);paragraph(note['recommendation'] if note else remediation)
        doc.add_heading('Comprobación de la mejora',3)
        paragraph(verification(item))
    doc.add_heading('6 Plan de mejora prioritario',1)
    if top:table(['Orden','Referencia','Acción','Resultado esperado'],[[str(i),o['id'],human_description(o)[2],'La comprobación específica acredita el control corregido'] for i,o in enumerate(top,1)])
    else:paragraph('No se priorizan correcciones con las condiciones disponibles. Completar las comprobaciones pendientes y conservar el seguimiento de los controles.')
    paragraph('El propietario asignará responsables y fechas según criticidad operativa y ventanas de cambio. Las señales pendientes deben validarse antes de decidir una remediación invasiva. Este plan no inventa compromisos de ejecución ni plazos acordados.')
    doc.add_heading('7 Ciberinteligencia y contexto público',1)
    intel=intelligence_context(data)
    paragraph('Se muestran los datos relevantes de cada fuente y su fecha de consulta. Una asociación pública puede estar desactualizada y no demuestra compromiso, titularidad o aplicabilidad de una CVE. La geolocalización de una IP describe al proveedor y su registro, no la ubicación física de la organización evaluada.')
    for row in intel['sources']:
        doc.add_heading(row['provider']+'  '+str(row['target']),2)
        label('Consulta',str(row.get('queried_at_utc',''))+' · '+str(row.get('status','')))
        if row.get('error'):paragraph(row['error']);continue
        if row['provider']=='VirusTotal':
            stats=row.get('analysis_stats',{})
            paragraph('Resultados disponibles: '+str(stats.get('malicious',0))+' maliciosos y '+str(stats.get('suspicious',0))+' sospechosos. Fecha del análisis del proveedor: '+str(row.get('analysis_date','no registrada'))+'. La fecha de consulta no equivale a una evaluación nueva. Cero detecciones no acredita ausencia de actividad maliciosa.')
            table(['Categoría','Resultados'],[[k,v] for k,v in stats.items()])
        else:
            for k,v in row.items():
                if k in ('provider','target','status','queried_at_utc','error') or v is None:continue
                label(k,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else str(v))
    if intel['correlations']:
        doc.add_heading('Correlación de vulnerabilidades públicas',2)
        for c in intel['correlations']:
            label(c['cve'],'Activos: '+', '.join(c.get('assets',[]))+' · fuentes: '+', '.join(c.get('sources',[])))
            paragraph('Presencia en el catálogo KEV consultado: '+('sí' if c.get('kev') else 'no')+'. '+c['status'])
            epss=(c.get('epss') or {}).get('data',{}).get('data',[])
            if epss:
                for e in epss:paragraph('EPSS: '+str(e.get('epss'))+' · percentil: '+str(e.get('percentile'))+' · fecha: '+str(e.get('date'))+'. Contextualiza probabilidad estimada de explotación del CVE; no confirma afectación del activo.')
            elif c.get('epss'):paragraph('EPSS: '+str(c['epss'].get('status'))+' · '+str(c['epss'].get('error','sin valor disponible')))

    doc.add_heading('8 Anexos de cobertura y trazabilidad',1)
    doc.add_heading('Reglas de puntuación y riesgo',2)
    paragraph(p['methodology'])
    paragraph('Riesgo cualitativo: impacto y probabilidad ordinal de 1 a 4. Producto 1–3 Bajo; 4–7 Medio; 8–11 Alto; 12–16 Crítico. Los indicios no validados se limitan a Alto. Informativo no es una vulnerabilidad ni una penalización nueva. Los criterios por condición están documentados en la metodología; esta matriz no utiliza CVSS.')
    doc.add_heading('Cobertura y límites',2);paragraph(n['limitations'])
    for warning in p['warnings']:paragraph(warning)
    table(['Herramienta','Resultado'],[[r.get('tool',''),r.get('status','')] for r in data.get('results',[])])
    if data.get('ip_candidates'):
        doc.add_heading('Descubrimiento y atribución de IPs públicas',2)
        table(['IP candidata','Fuentes','Relación y alcance'],[[r['ip'],', '.join(r['sources']),r['relation']+'; '+r['status']] for r in data['ip_candidates']])
        paragraph(data['ip_discovery']['limits'])
    long_lists=[o for o in data.get('observations',[]) if len(o.get('affected_urls',[]))>SHOWN_URLS]
    if long_lists:
        doc.add_heading('Ubicaciones afectadas por hallazgo',2)
        for o in long_lists:label(o['id']+' · '+o['title'],'\n'.join(o['affected_urls']))
    doc.add_heading('Comprobaciones por control',2)
    table(['Control y activo','Resultado','Crédito','Referencia'],[[c['control']+' · '+c['asset'],c['status'],str(round((1-c['penalty'])*100,2))+'%' if c['status']=='Evaluado' else 'Excluido' if c['status']=='No aplica' else '0%','Salida vinculada en JSON; '+c['source']] for c in p['checks']])
    for c in p['checks']:
        label(c['control']+' · '+c['asset'],c['reason'])
        if c.get('evidence'):paragraph('Soporte: '+redact(c['evidence']))
    tls=tls_summary(data)
    if tls:
        doc.add_heading('Protocolos TLS observados',2)
        cols=['Activo','TLS 1.0','TLS 1.1','TLS 1.2','TLS 1.3'];table(cols,[[r.get(k,'No observado') for k in cols] for r in tls])
    doc.add_heading('Registro y lectura de evidencias',2)
    paragraph('Los identificadores HAL relacionan las condiciones con las fuentes y capturas en el JSON de evidencias. Las capturas de navegador muestran la vista observada; las capturas de resultados muestran salidas registradas y no simulan una terminal en vivo. El hash permite contrastar integridad de un archivo, no demuestra por sí solo autenticidad o ausencia de manipulación previa.')
    for t in doc.tables:
        h=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(h)
        for row in t.rows:
            e=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(e)
    doc.core_properties.title='Informe de postura de seguridad externa'
    doc.core_properties.subject=org+' · '+data['domain']
    doc.core_properties.author=data.get('report_author') or 'Equipo evaluador'
    buf=BytesIO();doc.save(buf);return buf.getvalue()
